from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy import select, func
import asyncio
import json
import io
import csv

from app.database import get_db, AsyncSessionLocal
from app.models.academic import ProfessorSemester, Semester, Branch, Course, College, StudentEnrollment
from app.models.assignment import Assignment, AssignmentStatus
from app.models.submission import Submission, SubmissionStatus
from app.models.evaluation import EvaluationReport
from app.models.marks import MarksOverride
from app.models.report import MarksReport
from app.models.user import User
from app.dependencies import require_professor, get_current_user

router = APIRouter(prefix="/professor", tags=["professor"])


@router.get("/dashboard")
async def get_dashboard(
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(require_professor),
):
    ps_result = (await db.execute(
        select(ProfessorSemester).where(ProfessorSemester.professor_id == current_user.id)
    )).scalars().all()

    semesters_data = []
    for ps in ps_result:
        sem = await db.get(Semester, ps.semester_id)
        if not sem:
            continue
        branch = await db.get(Branch, sem.branch_id)
        course = await db.get(Course, branch.course_id) if branch else None
        college = await db.get(College, course.college_id) if course else None

        total_assignments = await db.scalar(
            select(func.count(Assignment.id)).where(
                Assignment.semester_id == sem.id,
                Assignment.status != AssignmentStatus.DELETED,
            )
        ) or 0
        total_students = await db.scalar(
            select(func.count(StudentEnrollment.id)).where(StudentEnrollment.semester_id == sem.id)
        ) or 0
        pending = await db.scalar(
            select(func.count(Submission.id))
            .join(Assignment, Assignment.id == Submission.assignment_id)
            .where(Assignment.semester_id == sem.id, Submission.status == SubmissionStatus.EVALUATING)
        ) or 0

        scores = (await db.execute(
            select(EvaluationReport.ai_score, Assignment.max_marks)
            .join(Submission, Submission.id == EvaluationReport.submission_id)
            .join(Assignment, Assignment.id == Submission.assignment_id)
            .where(Assignment.semester_id == sem.id, Submission.status == SubmissionStatus.EVALUATED)
        )).all()
        avg = round(sum(r[0] / r[1] * 100 for r in scores if r[1] > 0) / len(scores), 1) if scores else 0.0

        semesters_data.append({
            "semester": {
                "id": sem.id,
                "name": sem.name,
                "branch": branch.name if branch else "",
                "course": course.name if course else "",
                "college": college.name if college else "",
            },
            "analytics": {
                "total_assignments": total_assignments,
                "total_students": total_students,
                "average_marks_percentage": avg,
                "pending_evaluations": pending,
            },
        })
    return {
        "professor": {
            "id": current_user.id,
            "name": current_user.name,
            "professor_id": current_user.professor_id,
            "college_id": current_user.college_id,
        },
        "semesters": semesters_data,
    }


@router.get("/reports/{assignment_id}")
async def get_report(
    assignment_id: str,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(require_professor),
):
    a = await db.get(Assignment, assignment_id)
    if not a or a.created_by != current_user.id:
        raise HTTPException(status_code=403, detail="Access not permitted")

    report = (await db.execute(
        select(MarksReport).where(MarksReport.assignment_id == assignment_id)
        .order_by(MarksReport.generated_at.desc())
    )).scalar_one_or_none()
    if not report:
        raise HTTPException(status_code=404, detail="Report not yet generated")

    return {
        "id": report.id,
        "assignment_id": report.assignment_id,
        "pdf_url": report.pdf_url,
        "excel_url": report.excel_url,
        "csv_url": report.csv_url,
        "generated_at": report.generated_at.isoformat(),
    }


@router.get("/analytics/summary")
async def get_analytics_summary(
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(require_professor),
):
    """Aggregate analytics across all professor's semesters."""
    ps_result = (await db.execute(
        select(ProfessorSemester).where(ProfessorSemester.professor_id == current_user.id)
    )).scalars().all()

    semester_ids = [ps.semester_id for ps in ps_result]
    if not semester_ids:
        return {"total_assignments": 0, "total_students": 0, "total_submissions": 0,
                "total_evaluated": 0, "pending_evaluations": 0, "plagiarism_cases": 0,
                "average_score_percentage": 0.0}

    total_assignments = await db.scalar(
        select(func.count(Assignment.id)).where(
            Assignment.semester_id.in_(semester_ids),
            Assignment.status != AssignmentStatus.DELETED,
        )
    ) or 0

    total_students = await db.scalar(
        select(func.count(func.distinct(StudentEnrollment.student_id))).where(
            StudentEnrollment.semester_id.in_(semester_ids)
        )
    ) or 0

    total_submissions = await db.scalar(
        select(func.count(Submission.id))
        .join(Assignment, Assignment.id == Submission.assignment_id)
        .where(Assignment.semester_id.in_(semester_ids))
    ) or 0

    total_evaluated = await db.scalar(
        select(func.count(Submission.id))
        .join(Assignment, Assignment.id == Submission.assignment_id)
        .where(Assignment.semester_id.in_(semester_ids), Submission.status == SubmissionStatus.EVALUATED)
    ) or 0

    pending = await db.scalar(
        select(func.count(Submission.id))
        .join(Assignment, Assignment.id == Submission.assignment_id)
        .where(Assignment.semester_id.in_(semester_ids), Submission.status == SubmissionStatus.EVALUATING)
    ) or 0

    plagiarism = await db.scalar(
        select(func.count(Submission.id))
        .join(Assignment, Assignment.id == Submission.assignment_id)
        .where(Assignment.semester_id.in_(semester_ids), Submission.status == SubmissionStatus.SIMILARITY_REVIEW)
    ) or 0

    scores = (await db.execute(
        select(EvaluationReport.ai_score, Assignment.max_marks)
        .join(Submission, Submission.id == EvaluationReport.submission_id)
        .join(Assignment, Assignment.id == Submission.assignment_id)
        .where(Assignment.semester_id.in_(semester_ids))
    )).all()
    avg = round(sum(r[0] / r[1] * 100 for r in scores if r[1] > 0) / len(scores), 1) if scores else 0.0

    return {
        "total_assignments": total_assignments,
        "total_students": total_students,
        "total_submissions": total_submissions,
        "total_evaluated": total_evaluated,
        "pending_evaluations": pending,
        "plagiarism_cases": plagiarism,
        "average_score_percentage": avg,
    }


@router.post("/reports/{assignment_id}/generate")
async def trigger_report(
    assignment_id: str,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(require_professor),
):
    a = await db.get(Assignment, assignment_id)
    if not a or a.created_by != current_user.id:
        raise HTTPException(status_code=403, detail="Access not permitted")
    from app.tasks.report_tasks import generate_marks_report_task
    from app.tasks.evaluation_tasks import safe_delay
    safe_delay(generate_marks_report_task, str(assignment_id))
    return {"message": "Report generation triggered"}


# ── SSE: real-time updates for professor ─────────────────────────────────────

@router.get("/live/{assignment_id}")
async def professor_live(
    assignment_id: str,
    current_user: User = Depends(require_professor),
):
    """SSE endpoint — pushes submission stats every 10s so professor UI auto-refreshes."""
    async def event_generator():
        try:
            while True:
                async with AsyncSessionLocal() as db:
                    subs = (await db.execute(
                        select(Submission).where(Submission.assignment_id == assignment_id)
                    )).scalars().all()
                    total = len(subs)
                    evaluated = sum(1 for s in subs if s.status == SubmissionStatus.EVALUATED)
                    pending = sum(1 for s in subs if s.status == SubmissionStatus.EVALUATING)
                    flagged = sum(1 for s in subs if s.status == SubmissionStatus.SIMILARITY_REVIEW)
                    payload = json.dumps({
                        "type": "submission_update",
                        "total": total, "evaluated": evaluated,
                        "pending": pending, "flagged": flagged,
                    })
                yield f"data: {payload}\n\n"
                await asyncio.sleep(10)
        except asyncio.CancelledError:
            return

    return StreamingResponse(
        event_generator(),
        media_type="text/event-stream",
        headers={"Cache-Control": "no-cache", "X-Accel-Buffering": "no"},
    )


# ── Export submissions as CSV / Excel ─────────────────────────────────────────

@router.get("/assignments/{assignment_id}/export/csv")
async def export_csv(
    assignment_id: str,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(require_professor),
):
    rows = await _build_export_rows(assignment_id, current_user.id, db)
    buf = io.StringIO()
    writer = csv.DictWriter(buf, fieldnames=list(rows[0].keys()) if rows else [])
    writer.writeheader()
    writer.writerows(rows)
    return StreamingResponse(
        iter([buf.getvalue().encode("utf-8")]),
        media_type="text/csv",
        headers={"Content-Disposition": f"attachment; filename=report_{assignment_id[:8]}.csv"},
    )


@router.get("/assignments/{assignment_id}/export/excel")
async def export_excel(
    assignment_id: str,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(require_professor),
):
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    rows = await _build_export_rows(assignment_id, current_user.id, db)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Submissions"
    if rows:
        headers = list(rows[0].keys())
        ws.append(headers)
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="2563EB")
        for row in rows:
            ws.append([row[h] for h in headers])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        iter([buf.getvalue()]),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=report_{assignment_id[:8]}.xlsx"},
    )


async def _build_export_rows(assignment_id: str, professor_id: str, db) -> list[dict]:
    a = await db.get(Assignment, assignment_id)
    if not a or a.created_by != professor_id:
        raise HTTPException(status_code=403, detail="Access not permitted")

    enrollments = (await db.execute(
        select(StudentEnrollment).where(StudentEnrollment.semester_id == a.semester_id)
    )).scalars().all()

    rows = []
    for enr in enrollments:
        student = await db.get(User, enr.student_id)
        if not student:
            continue
        sub = (await db.execute(
            select(Submission).where(
                Submission.assignment_id == a.id,
                Submission.student_id == student.id,
                Submission.status != SubmissionStatus.REJECTED,
            ).order_by(Submission.submitted_at.desc())
        )).scalar_one_or_none()

        if not sub:
            rows.append({
                "Student": student.name, "Roll No": student.roll_number or "N/A",
                "Email": student.email, "Status": "Not Submitted",
                "Marks": "", "Total Marks": a.max_marks, "Percentage": "", "Grade": "",
                "Strengths": "", "Weaknesses": "", "Missing Points": "", "Suggestions": "",
                "Overall Remarks": "", "Plagiarism %": "", "Plagiarism Status": "", "Submitted At": "",
            })
            continue

        eval_r = (await db.execute(
            select(EvaluationReport).where(EvaluationReport.submission_id == sub.id)
        )).scalar_one_or_none()
        override = (await db.execute(
            select(MarksOverride).where(MarksOverride.submission_id == sub.id)
        )).scalar_one_or_none()

        final_score = override.revised_score if override else (eval_r.ai_score if eval_r else None)
        pct = round(final_score / a.max_marks * 100, 1) if final_score is not None and a.max_marks > 0 else None
        grade = eval_r.grade if eval_r else ""
        sim_pct = round(sub.similarity_score * 100, 1) if sub.similarity_score else ""
        sim_status = (
            "Exact Copy" if sub.similarity_score and sub.similarity_score >= 0.99 else
            "High Risk" if sub.similarity_score and sub.similarity_score >= 0.7 else
            "Warning" if sub.similarity_score and sub.similarity_score >= 0.4 else
            "Safe"
        ) if sub.similarity_score else "Safe"

        rows.append({
            "Student": student.name, "Roll No": student.roll_number or "N/A",
            "Email": student.email, "Status": sub.status.value,
            "Marks": final_score, "Total Marks": a.max_marks,
            "Percentage": f"{pct}%" if pct is not None else "",
            "Grade": grade,
            "Strengths": eval_r.strengths if eval_r else "",
            "Weaknesses": eval_r.areas_of_improvement if eval_r else "",
            "Missing Points": eval_r.missing_points if eval_r else "",
            "Suggestions": eval_r.suggestions if eval_r else "",
            "Overall Remarks": eval_r.overall_feedback if eval_r else "",
            "Plagiarism %": sim_pct,
            "Plagiarism Status": sim_status,
            "Submitted At": sub.submitted_at.strftime("%Y-%m-%d %H:%M UTC"),
        })

    return rows
