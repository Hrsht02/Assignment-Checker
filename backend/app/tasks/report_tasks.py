"""
Celery tasks for report generation (PDF / Excel / CSV).
"""
import asyncio
import io
import csv
import json

from app.tasks.celery_app import celery_app


def _run_async(coro):
    loop = asyncio.new_event_loop()
    try:
        return loop.run_until_complete(coro)
    finally:
        loop.close()


@celery_app.task(name="generate_marks_report")
def generate_marks_report_task(assignment_id: str):
    _run_async(_generate_marks_report_async(assignment_id))


async def _generate_marks_report_async(assignment_id: str):
    from app.database import AsyncSessionLocal
    from app.models.assignment import Assignment
    from app.models.submission import Submission, SubmissionStatus
    from app.models.evaluation import EvaluationReport
    from app.models.marks import MarksOverride
    from app.models.academic import StudentEnrollment   # ← fixed: was SectionEnrollment
    from app.models.user import User
    from app.models.report import MarksReport
    from app.models.notification import NotificationType
    from app.services import notification as notif_svc
    from app.services.email import send_report_email
    from app.services.storage import upload_bytes
    from sqlalchemy import select

    async with AsyncSessionLocal() as db:
        assignment = await db.get(Assignment, assignment_id)
        if not assignment:
            return

        # Get all students enrolled in this assignment's semester
        enrollments = (await db.execute(
            select(StudentEnrollment).where(StudentEnrollment.semester_id == assignment.semester_id)
        )).scalars().all()

        # Build report rows
        rows = []
        for enr in enrollments:
            student = await db.get(User, enr.student_id)
            if not student:
                continue

            sub = (await db.execute(
                select(Submission).where(
                    Submission.assignment_id == assignment.id,
                    Submission.student_id == student.id,
                    Submission.status != SubmissionStatus.REJECTED,
                ).order_by(Submission.submitted_at.desc())
            )).scalar_one_or_none()

            if not sub:
                rows.append({
                    "name": student.name,
                    "roll_number": student.roll_number or "N/A",
                    "email": student.email,
                    "marks": "Not Submitted",
                    "percentage": "—",
                    "grade": "—",
                    "submission_time": "N/A",
                    "status": "Not Submitted",
                })
                continue

            eval_r = (await db.execute(
                select(EvaluationReport).where(EvaluationReport.submission_id == sub.id)
            )).scalar_one_or_none()
            override = (await db.execute(
                select(MarksOverride).where(MarksOverride.submission_id == sub.id)
            )).scalar_one_or_none()

            if override:
                marks = str(override.revised_score)
                pct = f"{override.revised_score / assignment.max_marks * 100:.1f}%" if assignment.max_marks > 0 else "—"
            elif eval_r:
                marks = str(eval_r.ai_score)
                pct = f"{eval_r.percentage:.1f}%"
            else:
                marks = "Pending"
                pct = "—"

            grade = eval_r.grade if eval_r else "—"
            if override and assignment.max_marks > 0:
                pct_val = override.revised_score / assignment.max_marks * 100
                if pct_val >= 90: grade = "A+"
                elif pct_val >= 80: grade = "A"
                elif pct_val >= 70: grade = "B+"
                elif pct_val >= 60: grade = "B"
                elif pct_val >= 50: grade = "C"
                elif pct_val >= 40: grade = "D"
                else: grade = "F"

            rows.append({
                "name": student.name,
                "roll_number": student.roll_number or "N/A",
                "email": student.email,
                "marks": marks,
                "percentage": pct,
                "grade": grade,
                "submission_time": sub.submitted_at.strftime("%Y-%m-%d %H:%M UTC"),
                "status": sub.status.value,
            })

        # ── CSV ────────────────────────────────────────────────────────────
        csv_buf = io.StringIO()
        writer = csv.DictWriter(csv_buf, fieldnames=["name", "roll_number", "email", "marks", "percentage", "grade", "submission_time", "status"])
        writer.writeheader()
        writer.writerows(rows)
        csv_bytes = csv_buf.getvalue().encode("utf-8")

        # ── Excel ──────────────────────────────────────────────────────────
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Marks Report"
        headers = ["Student Name", "Roll No.", "Email", "Marks", "Percentage", "Grade", "Submission Time", "Status"]
        ws.append(headers)
        for r in rows:
            ws.append([r["name"], r["roll_number"], r["email"], r["marks"], r["percentage"], r["grade"], r["submission_time"], r["status"]])
        # Style header row
        from openpyxl.styles import Font, PatternFill
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="2563EB")
        excel_buf = io.BytesIO()
        wb.save(excel_buf)
        excel_bytes = excel_buf.getvalue()

        # ── PDF ────────────────────────────────────────────────────────────
        from reportlab.lib.pagesizes import landscape, letter
        from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
        from reportlab.lib import colors
        from reportlab.lib.styles import getSampleStyleSheet

        pdf_buf = io.BytesIO()
        doc = SimpleDocTemplate(pdf_buf, pagesize=landscape(letter))
        styles = getSampleStyleSheet()
        elements = [
            Paragraph(f"Marks Report: {assignment.title}", styles["Title"]),
            Spacer(1, 12),
        ]
        table_data = [headers] + [
            [r["name"], r["roll_number"], r["email"], r["marks"],
             r["percentage"], r["grade"], r["submission_time"], r["status"]]
            for r in rows
        ]
        table = Table(table_data, repeatRows=1)
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2563EB")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.whitesmoke),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, 0), 10),
            ("ALIGN", (0, 0), (-1, -1), "CENTER"),
            ("BACKGROUND", (0, 1), (-1, -1), colors.beige),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F0F4FF")]),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
            ("FONTSIZE", (0, 1), (-1, -1), 8),
        ]))
        elements.append(table)
        doc.build(elements)
        pdf_bytes = pdf_buf.getvalue()

        # ── Upload to storage ──────────────────────────────────────────────
        base = f"reports/{assignment_id}"
        max_retries = 3
        pdf_key = pdf_url = excel_key = excel_url = csv_key = csv_url = None

        for attempt in range(max_retries):
            try:
                pdf_key,   pdf_url   = await upload_bytes(pdf_bytes,   f"{base}/report.pdf",  "application/pdf")
                excel_key, excel_url = await upload_bytes(excel_bytes, f"{base}/report.xlsx",
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
                csv_key,   csv_url   = await upload_bytes(csv_bytes,   f"{base}/report.csv",  "text/csv")
                break
            except RuntimeError:
                if attempt == max_retries - 1:
                    return   # All retries failed — don't mark as available

        # ── Save report record ─────────────────────────────────────────────
        existing = (await db.execute(
            select(MarksReport).where(MarksReport.assignment_id == assignment.id)
        )).scalar_one_or_none()

        if existing:
            existing.pdf_url = pdf_url; existing.excel_url = excel_url; existing.csv_url = csv_url
            existing.pdf_key = pdf_key; existing.excel_key = excel_key; existing.csv_key = csv_key
        else:
            db.add(MarksReport(
                assignment_id=assignment.id,
                pdf_url=pdf_url, excel_url=excel_url, csv_url=csv_url,
                pdf_key=pdf_key, excel_key=excel_key, csv_key=csv_key,
            ))
        await db.commit()

        # ── Notify professor + admins ──────────────────────────────────────
        professor = await db.get(User, assignment.created_by)
        if professor:
            await notif_svc.create_notification(
                db, professor.id, NotificationType.REPORT_GENERATED,
                "Marks Report Ready",
                f"The marks report for '{assignment.title}' is available.",
                reference_id=assignment.id, reference_type="assignment",
            )
            await send_report_email(professor.email, professor.name, assignment.title)

        from app.models.user import UserRole
        admins = (await db.execute(
            select(User).where(User.role == UserRole.COLLEGE_ADMIN, User.college_id == professor.college_id if professor else None)
        )).scalars().all() if professor else []
        for admin in admins:
            await notif_svc.create_notification(
                db, admin.id, NotificationType.REPORT_GENERATED,
                "Marks Report Ready",
                f"Marks report for '{assignment.title}' has been generated.",
                reference_id=assignment.id, reference_type="assignment",
            )

        await db.commit()
